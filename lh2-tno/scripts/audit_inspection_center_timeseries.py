"""Audit inspection-centre XLSX evidence without promoting it to QRA input.

This tool records file hashes, workbook/sheet structure, signal-name classes,
and aggregate numeric ranges.  It never copies rows or creates a provider
packet.  Engineering units are accepted only through an explicit reviewed
JSON map; without that map the state-snapshot and frequency gates stay closed.
"""
from __future__ import annotations

import argparse
import hashlib
import json
import math
import sys
from collections import Counter
from pathlib import Path
from typing import Any

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "src"))

from lh2_tno.inspection_timeseries import (  # noqa: E402
    classify_signal,
    normalise_unit_map,
    resolve_unit,
    summarise_numeric,
)


SCHEMA = "prism.inspection_center_timeseries_audit.v1"


def sha256_file(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for block in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest()


def _json_load_unit_map(path: Path | None) -> dict[str, str]:
    if path is None:
        return {}
    payload = json.loads(path.read_text(encoding="utf-8"))
    if not isinstance(payload, dict):
        raise ValueError("unit map must be a JSON object")
    return normalise_unit_map(payload)


def _scenario_group(relative_path: str) -> str:
    if "인수" in relative_path:
        return "receipt"
    if "벤트" in relative_path or "vent" in relative_path.lower():
        return "vent"
    if "공급" in relative_path:
        return "supply"
    if "원본" in relative_path or "학습" in relative_path:
        return "source_or_training"
    return "unclassified"


def _empty_sheet(relative_path: str, name: str, error: str | None = None) -> dict[str, Any]:
    result: dict[str, Any] = {
        "name": name,
        "relative_path": relative_path,
        "header_row_index": None,
        "headers": [],
        "row_count": 0,
        "nonempty_row_count": 0,
        "signals": [],
        "unit_basis_status": "MISSING_OR_UNCONFIRMED",
        "numeric_columns": {},
    }
    if error:
        result["error"] = error
    return result


def audit_workbook(path: Path, root: Path, unit_map: dict[str, str]) -> dict[str, Any]:
    relative_path = path.relative_to(root).as_posix()
    record: dict[str, Any] = {
        "relative_path": relative_path,
        "scenario_group": _scenario_group(relative_path),
        "sha256": sha256_file(path),
        "size_bytes": path.stat().st_size,
        "status": "READ",
        "sheet_count": 0,
        "sheets": [],
    }
    try:
        from openpyxl import load_workbook
    except Exception as exc:  # pragma: no cover - environment-specific
        record["status"] = "READER_UNAVAILABLE"
        record["error"] = f"openpyxl import failed: {exc}"
        return record
    try:
        workbook = load_workbook(path, read_only=True, data_only=True)
    except Exception as exc:
        record["status"] = "WORKBOOK_READ_ERROR"
        record["error"] = f"{type(exc).__name__}: {exc}"
        return record

    try:
        names = list(workbook.sheetnames)
        record["sheet_count"] = len(names)
        for name in names:
            try:
                worksheet = workbook[name]
                rows = worksheet.iter_rows(values_only=True)
                header_row: tuple[Any, ...] | None = None
                header_index: int | None = None
                row_count = 0
                nonempty_row_count = 0
                for index, row in enumerate(rows, start=1):
                    if header_row is None:
                        if any(value not in (None, "") for value in row):
                            header_row = tuple(row)
                            header_index = index
                        continue
                    row_count += 1
                    if any(value not in (None, "") for value in row):
                        nonempty_row_count += 1
                if header_row is None:
                    record["sheets"].append(_empty_sheet(relative_path, name))
                    continue
                headers = [str(value).strip() if value not in (None, "") else f"column_{i + 1}" for i, value in enumerate(header_row)]
                classifications = [classify_signal(header) for header in headers]
                signals: list[dict[str, Any]] = []
                numeric_columns: dict[str, Any] = {}
                for header, classification in zip(headers, classifications):
                    unit, unit_status = resolve_unit(header, unit_map)
                    signal = {
                        "header": header,
                        "category": classification["category"],
                        "classification_basis": classification["basis"],
                        "unit": unit,
                        "unit_basis_status": unit_status,
                    }
                    signals.append(signal)
                    if classification["category"] not in {"time", "sequence_metadata", "control_or_valve_state"}:
                        numeric_columns[header] = {
                            "category": classification["category"],
                            "unit": unit,
                            "unit_basis_status": unit_status,
                            "numeric_count": 0,
                            "min": None,
                            "max": None,
                        }
                # A second pass is intentional: it preserves header detection
                # while keeping the first pass simple and avoiding raw rows in
                # the output.  Malformed sheets are captured at sheet scope.
                try:
                    rows_again = worksheet.iter_rows(values_only=True)
                    for index, row in enumerate(rows_again, start=1):
                        if header_index is None or index <= header_index:
                            continue
                        for column_index, header in enumerate(headers):
                            value = row[column_index] if column_index < len(row) else None
                            if header not in numeric_columns:
                                continue
                            if isinstance(value, (int, float)) and not isinstance(value, bool) and math.isfinite(float(value)):
                                summary = numeric_columns[header]
                                summary["numeric_count"] += 1
                                number = float(value)
                                summary["min"] = number if summary["min"] is None else min(summary["min"], number)
                                summary["max"] = number if summary["max"] is None else max(summary["max"], number)
                except Exception as exc:
                    record["sheets"].append(_empty_sheet(relative_path, name, f"row scan failed: {type(exc).__name__}: {exc}"))
                    continue
                record["sheets"].append({
                    "name": name,
                    "relative_path": relative_path,
                    "header_row_index": header_index,
                    "headers": headers,
                    "row_count": row_count,
                    "nonempty_row_count": nonempty_row_count,
                    "signals": signals,
                    "unit_basis_status": "CONFIRMED_BY_REVIEWED_MAP" if unit_map and all(item["unit_basis_status"] == "CONFIRMED_BY_REVIEWED_MAP" for item in signals if item["category"] not in {"time", "sequence_metadata", "control_or_valve_state"}) else "MISSING_OR_UNCONFIRMED",
                    "numeric_columns": numeric_columns,
                })
            except Exception as exc:
                record["sheets"].append(_empty_sheet(relative_path, name, f"sheet scan failed: {type(exc).__name__}: {exc}"))
    finally:
        try:
            workbook.close()
        except Exception:
            pass
    if any("error" in sheet for sheet in record["sheets"]):
        record["status"] = "READ_WITH_SHEET_ERRORS"
    return record


def _aggregate(workbooks: list[dict[str, Any]], unit_map: dict[str, str]) -> dict[str, Any]:
    categories: Counter[str] = Counter()
    unit_statuses: Counter[str] = Counter()
    error_count = 0
    sheet_count = 0
    for workbook in workbooks:
        if workbook.get("status") in {"WORKBOOK_READ_ERROR", "READER_UNAVAILABLE"}:
            error_count += 1
        for sheet in workbook.get("sheets", []):
            sheet_count += 1
            for signal in sheet.get("signals", []):
                categories[signal["category"]] += 1
            # Count only columns that actually contained at least one numeric
            # value.  A text/action column with an unknown header is not a
            # missing-unit engineering signal.
            for numeric in sheet.get("numeric_columns", {}).values():
                if numeric.get("numeric_count", 0) > 0:
                    unit_statuses[numeric["unit_basis_status"]] += 1
            if "error" in sheet:
                error_count += 1
    return {
        "workbook_count": len(workbooks),
        "sheet_count": sheet_count,
        "workbook_or_sheet_error_count": error_count,
        "signal_category_count": dict(sorted(categories.items())),
        "numeric_signal_unit_status_count": dict(sorted(unit_statuses.items())),
        "reviewed_unit_map_supplied": bool(unit_map),
    }


def audit_root(source_root: Path, unit_map: dict[str, str]) -> dict[str, Any]:
    source_root = source_root.resolve()
    files = sorted(path for path in source_root.rglob("*.xlsx") if path.is_file() and not path.name.startswith("~$"))
    workbooks = [audit_workbook(path, source_root, unit_map) for path in files]
    aggregate = _aggregate(workbooks, unit_map)
    return {
        "schema_version": SCHEMA,
        "status": "SCREENING_ONLY_UNIT_BASIS_UNCONFIRMED" if not unit_map else "SCREENING_ONLY_REVIEWED_UNIT_MAP_NOT_STATE_BOUND",
        "source_root": str(source_root),
        "source_file_count": len(files),
        "aggregate": aggregate,
        "files": workbooks,
        "promotion_gates": {
            "state_snapshot_binding_allowed": False,
            "frequency_applied": False,
            "sensor_frequency_learning_used": False,
            "provider_packet_created": False,
            "facility_qra_claim": False,
            "reason": "XLSX time-series evidence is an audit input only; an exact reviewed unit map, tag semantics, calibration/quality evidence, and an explicit state-snapshot adapter are still required.",
        },
    }


def public_summary(report: dict[str, Any]) -> dict[str, Any]:
    """Remove local paths and raw header detail from the tracked summary."""

    summary = {
        "schema_version": report["schema_version"],
        "status": report["status"],
        "source_root": "<inspection-center-data-root>",
        "source_file_count": report["source_file_count"],
        "aggregate": report["aggregate"],
        "promotion_gates": report["promotion_gates"],
        "files": [],
    }
    for workbook in report.get("files", []):
        item = {
            "relative_path": workbook.get("relative_path"),
            "scenario_group": workbook.get("scenario_group"),
            "sha256": workbook.get("sha256"),
            "size_bytes": workbook.get("size_bytes"),
            "status": workbook.get("status"),
            "sheet_count": workbook.get("sheet_count", 0),
            "error": workbook.get("error"),
            "sheets": [],
        }
        for sheet in workbook.get("sheets", []):
            category_count = Counter(signal.get("category", "unknown") for signal in sheet.get("signals", []))
            numeric_columns = sheet.get("numeric_columns", {})
            item["sheets"].append({
                "name": sheet.get("name"),
                "header_row_index": sheet.get("header_row_index"),
                "row_count": sheet.get("row_count", 0),
                "nonempty_row_count": sheet.get("nonempty_row_count", 0),
                "signal_category_count": dict(sorted(category_count.items())),
                "numeric_candidate_column_count": sum(1 for value in numeric_columns.values() if value.get("numeric_count", 0) > 0),
                "unit_unconfirmed_numeric_column_count": sum(1 for value in numeric_columns.values() if value.get("numeric_count", 0) > 0 and value.get("unit_basis_status") != "CONFIRMED_BY_REVIEWED_MAP"),
                "error": sheet.get("error"),
            })
        summary["files"].append(item)
    return summary


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("--source-root", type=Path, required=True)
    parser.add_argument("--output", type=Path, required=True)
    parser.add_argument("--summary-output", type=Path)
    parser.add_argument("--unit-map", type=Path)
    args = parser.parse_args()
    try:
        unit_map = _json_load_unit_map(args.unit_map)
        report = audit_root(args.source_root, unit_map)
    except Exception as exc:
        print(json.dumps({"status": "AUDIT_ERROR", "error": f"{type(exc).__name__}: {exc}"}, ensure_ascii=False))
        return 2
    args.output.parent.mkdir(parents=True, exist_ok=True)
    args.output.write_text(json.dumps(report, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    if args.summary_output:
        args.summary_output.parent.mkdir(parents=True, exist_ok=True)
        args.summary_output.write_text(json.dumps(public_summary(report), ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(json.dumps({"status": report["status"], "output": str(args.output), "aggregate": report["aggregate"]}, ensure_ascii=False))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
