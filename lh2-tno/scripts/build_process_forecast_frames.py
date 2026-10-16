"""Build PRISM forecast frames from an Excel process-model export.

This command is a conservative digital-twin input adapter.  The user must
provide the component, facility and operating mode, plus an explicit mapping
from workbook columns to engineering quantities.  The command does not infer
operating mode, accident frequency, release duration, or facility risk.
"""
from __future__ import annotations

import argparse
import hashlib
import json
from pathlib import Path
from typing import Any

from lh2_tno.process_forecast import build_forecast_frames


def _sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for block in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest()


def _load_mapping(path: Path) -> dict[str, Any]:
    value = json.loads(path.read_text(encoding="utf-8"))
    if not isinstance(value, dict):
        raise ValueError("mapping JSON must contain an object")
    return value


def _rows_from_workbook(path: Path, *, sheet: str | None, sheet_index: int) -> list[dict[str, Any]]:
    try:
        import openpyxl
    except ImportError as error:  # pragma: no cover - exercised in user environment
        raise RuntimeError("openpyxl is required; install prism-risk[xlsx]") from error
    workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
    if sheet is not None:
        if sheet not in workbook.sheetnames:
            raise ValueError(f"sheet {sheet!r} is missing; available={workbook.sheetnames}")
        worksheet = workbook[sheet]
    else:
        try:
            worksheet = workbook.worksheets[sheet_index]
        except IndexError as error:
            raise ValueError(f"sheet-index {sheet_index} is outside workbook") from error
    iterator = worksheet.iter_rows(values_only=True)
    try:
        headers = [str(value).strip() if value is not None else "" for value in next(iterator)]
    except StopIteration as error:
        raise ValueError("worksheet has no header row") from error
    if not all(headers):
        raise ValueError("worksheet header row contains an empty column")
    rows: list[dict[str, Any]] = []
    for raw in iterator:
        if all(value is None for value in raw):
            continue
        rows.append(dict(zip(headers, raw, strict=False)))
    if not rows:
        raise ValueError("worksheet has no data rows")
    return rows


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--source", type=Path, required=True)
    parser.add_argument("--mapping", type=Path, required=True, help="JSON mapping keyed by PRISM tag")
    parser.add_argument("--output", type=Path, required=True)
    parser.add_argument("--facility-id", required=True)
    parser.add_argument("--component-id", required=True)
    parser.add_argument("--operating-mode", required=True)
    parser.add_argument("--source-id", default="process-workbook")
    parser.add_argument("--sheet")
    parser.add_argument("--sheet-index", type=int, default=0)
    parser.add_argument("--time-column")
    parser.add_argument("--sample-period-s", type=float, default=1.0)
    parser.add_argument("--max-rows", type=int)
    args = parser.parse_args()
    rows = _rows_from_workbook(args.source, sheet=args.sheet, sheet_index=args.sheet_index)
    mapping = _load_mapping(args.mapping)
    record = build_forecast_frames(
        rows,
        facility_id=args.facility_id,
        component_id=args.component_id,
        operating_mode=args.operating_mode,
        mapping=mapping,
        source_id=args.source_id,
        source_provenance=args.source.as_posix(),
        time_column=args.time_column,
        sample_period_s=args.sample_period_s,
        source_digest=_sha256(args.source),
        max_rows=args.max_rows,
    )
    args.output.parent.mkdir(parents=True, exist_ok=True)
    args.output.write_text(json.dumps(record, ensure_ascii=False, indent=2, default=str) + "\n", encoding="utf-8")
    print(json.dumps({
        "status": record["status"],
        "schema": record["schema_version"],
        "frame_count": len(record["frames"]),
        "output": str(args.output),
    }, ensure_ascii=False))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
