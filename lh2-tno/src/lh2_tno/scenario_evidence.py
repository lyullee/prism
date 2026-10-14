"""Read-only audit for scenario registers and time-series exports.

The inspection-centre workbooks and CSV registers are useful evidence for
equipment states, operating steps, and accident-family candidates.  They are
not resolved provider release packets.  This module inventories them without
turning a pressure/temperature trace, a scenario description, or a vent
command into a release rate or accident frequency.
"""
from __future__ import annotations

import csv
import hashlib
import json
import re
from collections import Counter
from collections.abc import Mapping
from datetime import date, datetime, time
from pathlib import Path
from statistics import median
from typing import Any


SCENARIO_EVIDENCE_SCHEMA = "prism.scenario_evidence_audit.v1"
PROVIDER_PACKET_REQUEST_CATALOG_SCHEMA = "prism.provider_packet_request_catalog.v1"
PROVIDER_RELEASE_REQUIRED_FIELDS = (
    "provider_export_schema",
    "boundary_kind",
    "event_kind",
    "event_id",
    "component_id",
    "port_id",
    "position_enu_m",
    "direction_enu",
    "steps",
    "duration_s",
    "available_mass_kg",
    "cumulative_mass_out_kg",
    "termination_provenance",
    "property_basis",
    "provider_source_digest",
    "provider_result_digest",
    # Required for accidental-leak frequency exact matching.  A dedicated
    # emergency-relief packet intentionally omits this random-opening field.
    "failure_opening_diameter_m",
)

_FIELD_ALIASES = {
    "scenario": ("시나리오", "scenario", "scenario_id"),
    "operating_mode": ("운전모드", "운전 모드", "operating_mode", "operation_mode", "mode"),
    "asset": ("탱크", "tank", "asset_id", "component_id", "equipment"),
    "step": ("절차번호", "step", "sequence", "step_id"),
    "operation": ("동작내용", "action", "operation", "command"),
    "hazard_class": ("위험성분류", "hazard_class", "hazard_type"),
    "leak_failure": ("누출_예상고장", "leak_failure", "release_failure"),
    "leak_hazard": ("누출_위험", "leak_hazard", "release_hazard"),
    "fire_hazard": ("화재_위험", "fire_hazard"),
    "explosion_hazard": ("폭발_위험", "explosion_hazard"),
}

_OPERATING_MODE_LABELS: tuple[tuple[str, str], ...] = (
    ("인수", "receiving"),
    ("receiving", "receiving"),
    ("저장(벤트)", "storage_vent"),
    ("저장", "storage_vent"),
    ("벤트", "storage_vent"),
    ("vent", "storage_vent"),
    ("공급", "supply"),
    ("supply", "supply"),
    ("대기", "standby"),
    ("스텐바이", "standby"),
    ("standby", "standby"),
)


def _operating_mode_candidate(
    explicit_mode: Any,
    scenario: Any,
    relative_path: str,
) -> tuple[str | None, str]:
    """Map source labels to a HAZOP organization hint, never to live state.

    The live ``HazopStateFrame.operating_mode`` remains the authoritative
    input for scenario selection.  This helper only preserves a traceable
    candidate from a procedure register or folder name for provider requests.
    """
    explicit_texts = [str(explicit_mode).strip()] if explicit_mode else []
    explicit_matches = [
        (mode, token)
        for text in explicit_texts
        for token, mode in _OPERATING_MODE_LABELS
        if token.casefold() in text.casefold()
    ]
    if explicit_matches:
        explicit_modes = {mode for mode, _ in explicit_matches}
        if len(explicit_modes) != 1:
            return None, "ambiguous_source_label"
        return explicit_matches[0][0], f"procedure_field:{explicit_matches[0][1]}"

    texts = [str(value).strip() for value in (scenario, relative_path) if value]
    matches = [
        (mode, token)
        for text in texts
        for token, mode in _OPERATING_MODE_LABELS
        if token.casefold() in text.casefold()
    ]
    if not matches:
        return None, "unresolved_source_label"
    modes = {mode for mode, _ in matches}
    if len(modes) != 1:
        return None, "ambiguous_source_label"
    return matches[0][0], f"procedure_label:{matches[0][1]}"

_FAMILY_BY_HAZARD_CLASS = {
    "flange_connect": ("pipe_hole", "vessel_hole"),
    "purge_vent": ("vent_hole",),
    "relief_vent": ("vent_hole",),
    "valve_close": ("valve_hole",),
    "main_line_open": ("pipe_hole",),
    "supply_start": ("pipe_hole", "valve_hole"),
    "vaporizer_pass": ("vaporizer_hole",),
    "pressure_rise_monitor": ("vent_hole", "inventory_over_600s"),
}

_FAMILY_PROVIDER_REQUIREMENTS = {
    "inventory_over_600s": {
        "acceptable_operators": ("equipment_inventory_over_600s", "equipment_liquid_over_600s"),
        "termination_basis": "inventory_over_600s",
        "required_duration_s": 600.0,
    },
    "inventory_depletion": {
        "acceptable_operators": ("equipment_inventory_depletion", "equipment_liquid_release"),
        "termination_basis": "inventory_depletion",
    },
}


def _sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for block in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest()


def _decode_csv(path: Path) -> tuple[str, str]:
    raw = path.read_bytes()
    for encoding in ("utf-8-sig", "cp949", "euc-kr", "utf-16"):
        try:
            return raw.decode(encoding), encoding
        except UnicodeDecodeError:
            continue
    raise UnicodeDecodeError("unknown", raw, 0, min(1, len(raw)), "unsupported scenario CSV encoding")


def _value(row: dict[str, Any], aliases: tuple[str, ...]) -> Any:
    for key in aliases:
        if key in row and row[key] not in (None, ""):
            return row[key]
    return None


def _scenario_record(path: Path, root: Path, row_number: int, row: dict[str, Any]) -> dict[str, Any]:
    hazard_class = _value(row, _FIELD_ALIASES["hazard_class"])
    family = _FAMILY_BY_HAZARD_CLASS.get(str(hazard_class), ())
    scenario = _value(row, _FIELD_ALIASES["scenario"])
    operating_mode_candidate, operating_mode_source = _operating_mode_candidate(
        _value(row, _FIELD_ALIASES["operating_mode"]),
        scenario,
        str(path.relative_to(root)),
    )
    return {
        "file": str(path.relative_to(root)),
        "row_number": row_number,
        "scenario": scenario,
        "operating_mode_candidate": operating_mode_candidate,
        "operating_mode_source": operating_mode_source,
        "operating_mode_is_runtime_input": False,
        "asset": _value(row, _FIELD_ALIASES["asset"]),
        "step": _value(row, _FIELD_ALIASES["step"]),
        "operation": _value(row, _FIELD_ALIASES["operation"]),
        "hazard_class": hazard_class,
        "candidate_accident_families": list(family),
        "hazard_text": {
            "leak_failure": _value(row, _FIELD_ALIASES["leak_failure"]),
            "leak_hazard": _value(row, _FIELD_ALIASES["leak_hazard"]),
            "fire_hazard": _value(row, _FIELD_ALIASES["fire_hazard"]),
            "explosion_hazard": _value(row, _FIELD_ALIASES["explosion_hazard"]),
        },
        "promotion_status": "SCENARIO_ONLY_PROVIDER_PACKET_MISSING",
        "missing_provider_fields": list(PROVIDER_RELEASE_REQUIRED_FIELDS),
        "frequency_assigned": False,
    }


def _audit_csv(path: Path, root: Path) -> dict[str, Any]:
    text, encoding = _decode_csv(path)
    rows = list(csv.DictReader(text.splitlines()))
    records = []
    for index, row in enumerate(rows, start=2):
        if any(_value(row, aliases) is not None for aliases in _FIELD_ALIASES.values()):
            records.append(_scenario_record(path, root, index, row))
    return {
        "file": str(path.relative_to(root)),
        "sha256": _sha256(path),
        "encoding": encoding,
        "row_count": len(rows),
        "scenario_row_count": len(records),
        "scenario_records": records,
    }


def _classify_header(header: list[str]) -> str:
    joined = "|".join(header)
    if any(key in joined for key in ("timestamp", "일시", "시간")):
        return "timeseries"
    if any(key in joined for key in ("from_step", "progress_pct", "action", "절차", "동작")):
        return "sequence_timeseries"
    return "metadata_or_unknown"


def _time_axis_index(header: list[str]) -> int | None:
    """Return an explicit time column index without assigning units to numbers."""
    for index, value in enumerate(header):
        lowered = value.strip().lower()
        if any(token in lowered for token in ("timestamp", "date_time", "datetime", "일시", "시간")):
            return index
    return None


def _explicit_unit(header: str) -> str | None:
    match = re.search(r"\(([^()]*)\)", header)
    if not match:
        return None
    value = match.group(1).strip()
    return value or None


def _sensor_column_evidence(header: list[str], time_index: int | None) -> list[dict[str, Any]]:
    evidence = []
    for index, value in enumerate(header):
        if index == time_index or not value.strip():
            continue
        lowered = value.strip().lower()
        upper = value.strip().upper()
        sensor_token = any(token in lowered for token in (
            "sensor", "pressure", "temperature", "flow", "level", "압력", "온도", "유량", "레벨", "차압",
        ))
        short_tag = re.search(r"(?:^|[^A-Z])(?:DPT|PT|TT|FT|LT)\s*[-_ ]?\d", upper) is not None
        if sensor_token or short_tag:
            evidence.append({
                "column_index": index,
                "header": value,
                "explicit_unit": _explicit_unit(value),
            })
    return evidence


def _parse_time_value(value: Any) -> datetime | None:
    """Parse typed/ISO time values only; numeric samples stay untyped."""
    if isinstance(value, datetime):
        return value
    if isinstance(value, date):
        return datetime.combine(value, time.min)
    if isinstance(value, time):
        return datetime.combine(date(2000, 1, 1), value)
    if not isinstance(value, str):
        return None
    text = value.strip()
    if not text:
        return None
    for candidate in (text, text.replace("/", "-"), text.rstrip("Z")):
        try:
            return datetime.fromisoformat(candidate)
        except ValueError:
            continue
    for pattern in ("%H:%M:%S", "%H:%M"):
        try:
            return datetime.strptime(text, pattern).replace(year=2000, month=1, day=1)
        except ValueError:
            continue
    time_fragment = re.search(r"(?<!\d)(\d{1,2}:\d{2}(?::\d{2})?)(?!\d)", text)
    if time_fragment:
        for pattern in ("%H:%M:%S", "%H:%M"):
            try:
                return datetime.strptime(time_fragment.group(1), pattern).replace(
                    year=2000, month=1, day=1
                )
            except ValueError:
                continue
    return None


def _telemetry_evidence(header: list[str], data_rows: list[tuple[Any, ...]], kind: str) -> dict[str, Any]:
    time_index = _time_axis_index(header)
    sensor_columns = _sensor_column_evidence(header, time_index)
    parsed_times = []
    if time_index is not None:
        for row in data_rows:
            if time_index < len(row):
                parsed = _parse_time_value(row[time_index])
                if parsed is not None:
                    parsed_times.append(parsed)
    intervals = [
        (later - earlier).total_seconds()
        for earlier, later in zip(parsed_times, parsed_times[1:])
        if (later - earlier).total_seconds() >= 0
    ]
    metadata_presence = {
        "component_mapping": any(
            any(token in value.strip().lower() for token in ("component_id", "asset_id", "equipment", "설비", "탱크"))
            for value in header
        ),
        "sensor_location": any(
            any(token in value.strip().lower() for token in ("sensor_location", "location", "위치", "노즐"))
            for value in header
        ),
        "pressure_basis": any(
            any(token in value.strip().lower() for token in ("pressure_basis", "gauge", "absolute", "압력기준", "게이지", "절대"))
            for value in header
        ),
        "provider_release_ownership": any(
            any(token in value.strip().lower() for token in ("provider", "release_owner", "방출공급자"))
            for value in header
        ),
    }
    missing_metadata = [name for name, present in metadata_presence.items() if not present]
    candidate = time_index is not None and bool(sensor_columns)
    return {
        "kind": kind,
        "time_axis_header": header[time_index] if time_index is not None else None,
        "time_axis_present": time_index is not None,
        "sample_count": len(data_rows),
        "parseable_time_count": len(parsed_times),
        "first_time": parsed_times[0].isoformat() if parsed_times else None,
        "last_time": parsed_times[-1].isoformat() if parsed_times else None,
        "sample_interval_s": {
            "count": len(intervals),
            "min": min(intervals) if intervals else None,
            "median": median(intervals) if intervals else None,
            "max": max(intervals) if intervals else None,
            "numeric_time_values_withheld": time_index is not None and any(
                time_index < len(row) and isinstance(row[time_index], (int, float)) and not isinstance(row[time_index], bool)
                for row in data_rows
            ),
        },
        "sensor_columns": sensor_columns,
        "metadata_presence": metadata_presence,
        "missing_metadata": missing_metadata,
        "telemetry_candidate": candidate,
        "promotion_status": (
            "TELEMETRY_CANDIDATE_METADATA_INCOMPLETE" if candidate and missing_metadata
            else "TELEMETRY_CANDIDATE_READY_FOR_PROVIDER_MAPPING" if candidate
            else "NO_TELEMETRY_CANDIDATE"
        ),
    }


def _audit_workbook(path: Path, root: Path) -> dict[str, Any]:
    try:
        from openpyxl import load_workbook
    except ImportError:
        return {
            "file": str(path.relative_to(root)),
            "sha256": _sha256(path),
            "status": "WORKBOOK_READER_UNAVAILABLE",
            "sheet_count": None,
            "sheets": [],
        }
    workbook = load_workbook(path, read_only=True, data_only=True)
    sheets = []
    for worksheet in workbook.worksheets:
        iterator = worksheet.iter_rows(values_only=True)
        first = next(iterator, ())
        header = [str(value) if value is not None else "" for value in first]
        data_rows = list(iterator)
        row_count = 1 + len(data_rows)
        telemetry = _telemetry_evidence(header, data_rows, _classify_header(header))
        sheets.append({
            "sheet": worksheet.title,
            "row_count": row_count,
            "column_count": len(header),
            "header": header[:32],
            "kind": _classify_header(header),
            "promotion_status": "TIME_SERIES_OR_SEQUENCE_ONLY_PROVIDER_PACKET_MISSING",
            "telemetry": telemetry,
        })
    telemetry_candidate_sheet_count = sum(
        1 for sheet in sheets if sheet["telemetry"]["telemetry_candidate"]
    )
    telemetry_metadata_incomplete_sheet_count = sum(
        1
        for sheet in sheets
        if sheet["telemetry"]["promotion_status"] == "TELEMETRY_CANDIDATE_METADATA_INCOMPLETE"
    )
    return {
        "file": str(path.relative_to(root)),
        "sha256": _sha256(path),
        "status": "READ_ONLY_WORKBOOK_AUDITED",
        "sheet_count": len(sheets),
        "telemetry_candidate_sheet_count": telemetry_candidate_sheet_count,
        "telemetry_metadata_incomplete_sheet_count": telemetry_metadata_incomplete_sheet_count,
        "sheets": sheets,
    }


def audit_scenario_root(root: str | Path) -> dict[str, Any]:
    """Audit center-style CSV/XLSX scenario evidence without promotion."""
    root_path = Path(root)
    if not root_path.exists():
        return {
            "schema": SCENARIO_EVIDENCE_SCHEMA,
            "status": "MISSING_SOURCE_ROOT",
            "source_root": str(root_path),
            "blocking_reasons": ["scenario evidence root does not exist"],
        }
    csv_reports = []
    workbook_reports = []
    errors = []
    for path in sorted(root_path.rglob("*")):
        if not path.is_file() or any(part in {".git", "__pycache__", ".venv"} for part in path.parts):
            continue
        try:
            if path.suffix.lower() == ".csv":
                csv_reports.append(_audit_csv(path, root_path))
            elif path.suffix.lower() in {".xlsx", ".xlsm"}:
                workbook_reports.append(_audit_workbook(path, root_path))
        except (OSError, UnicodeError, ValueError, RuntimeError) as error:
            errors.append({"file": str(path.relative_to(root_path)), "error": str(error)})

    records = [record for report in csv_reports for record in report["scenario_records"]]
    family_counts = Counter(
        family
        for record in records
        for family in record["candidate_accident_families"]
    )
    sequence_sheet_count = sum(
        1
        for report in workbook_reports
        for sheet in report.get("sheets", [])
        if sheet.get("kind") == "sequence_timeseries"
    )
    timeseries_sheet_count = sum(
        1
        for report in workbook_reports
        for sheet in report.get("sheets", [])
        if sheet.get("kind") == "timeseries"
    )
    telemetry_candidate_sheet_count = sum(
        report.get("telemetry_candidate_sheet_count", 0) for report in workbook_reports
    )
    telemetry_metadata_incomplete_sheet_count = sum(
        report.get("telemetry_metadata_incomplete_sheet_count", 0) for report in workbook_reports
    )
    telemetry_parseable_time_sheet_count = sum(
        1
        for report in workbook_reports
        for sheet in report.get("sheets", [])
        if sheet.get("telemetry", {}).get("parseable_time_count", 0) > 1
    )
    interval_profiles = Counter()
    for report in workbook_reports:
        for sheet in report.get("sheets", []):
            telemetry = sheet.get("telemetry", {})
            median_interval = telemetry.get("sample_interval_s", {}).get("median")
            if median_interval is not None and telemetry.get("telemetry_candidate"):
                interval_profiles[str(median_interval)] += 1
    blockers = []
    if records:
        blockers.append("scenario registers describe candidate hazards but contain no resolved provider release packet")
    if sequence_sheet_count or timeseries_sheet_count:
        blockers.append("workbook sequence/time-series exports lack release mass-flow, termination, geometry, and provider ownership fields")
    if telemetry_candidate_sheet_count:
        blockers.append("workbook sensor traces have a time axis but lack component mapping, sensor location, pressure basis, and provider release ownership")
    if errors:
        blockers.append("one or more scenario evidence files could not be read")
    return {
        "schema": SCENARIO_EVIDENCE_SCHEMA,
        "status": "SCENARIO_EVIDENCE_ONLY" if records or workbook_reports else "NO_SCENARIO_EVIDENCE",
        "source_root": str(root_path),
        "files": {
            "csv_count": len(csv_reports),
            "workbook_count": len(workbook_reports),
            "csv_rows": sum(report["row_count"] for report in csv_reports),
            "scenario_rows": len(records),
            "sequence_sheet_count": sequence_sheet_count,
            "timeseries_sheet_count": timeseries_sheet_count,
            "telemetry_candidate_sheet_count": telemetry_candidate_sheet_count,
            "telemetry_metadata_incomplete_sheet_count": telemetry_metadata_incomplete_sheet_count,
            "telemetry_parseable_time_sheet_count": telemetry_parseable_time_sheet_count,
            "telemetry_median_interval_s_profiles": dict(sorted(interval_profiles.items(), key=lambda item: float(item[0]))),
        },
        "candidate_accident_family_counts": dict(sorted(family_counts.items())),
        "scenario_records": records,
        "csv_reports": csv_reports,
        "workbook_reports": workbook_reports,
        "provider_promotion": {
            "qualified_packet_count": 0,
            "frequency_assigned": False,
            "required_fields": list(PROVIDER_RELEASE_REQUIRED_FIELDS),
            "scenario_rows_are_not_release_packets": True,
        },
        "errors": errors,
        "blocking_reasons": list(dict.fromkeys(blockers)),
    }


def build_provider_packet_request_catalog(audit: Mapping[str, Any]) -> dict[str, Any]:
    """Turn scenario-only hazards into an explicit provider data request list.

    The result is a work queue, not a release catalog. It preserves candidate
    family and source-row provenance while withholding frequency and all
    release values until a validated ``prism.resolved_release.v1`` packet
    arrives. This keeps scenario registers actionable without promoting them to
    accident histories.
    """
    records = audit.get("scenario_records", [])
    if not isinstance(records, list):
        raise ValueError("scenario audit must contain a scenario_records list")
    requests = []
    family_counts: Counter[str] = Counter()
    for record in records:
        if not isinstance(record, Mapping):
            raise ValueError("scenario records must be objects")
        families = tuple(str(value) for value in record.get("candidate_accident_families", ()))
        if not families:
            continue
        identity = {
            "file": record.get("file"),
            "row_number": record.get("row_number"),
            "scenario": record.get("scenario"),
            "operating_mode_candidate": record.get("operating_mode_candidate"),
            "operating_mode_source": record.get("operating_mode_source"),
            "asset": record.get("asset"),
            "step": record.get("step"),
            "operation": record.get("operation"),
        }
        request_id = "scenario-" + hashlib.sha256(
            json.dumps(identity, ensure_ascii=False, sort_keys=True, separators=(",", ":")).encode("utf-8")
        ).hexdigest()[:16]
        for family in families:
            family_counts[family] += 1
        family_requirements = [
            {"family": family, **_FAMILY_PROVIDER_REQUIREMENTS[family]}
            for family in families
            if family in _FAMILY_PROVIDER_REQUIREMENTS
        ]
        requests.append({
            "request_id": request_id,
            "source": identity,
            "runtime_operating_mode_required": True,
            "candidate_accident_families": list(families),
            "hazard_class": record.get("hazard_class"),
            "provider_export_schema": "prism.resolved_release.v1",
            "required_provider_fields": list(
                record.get("missing_provider_fields") or PROVIDER_RELEASE_REQUIRED_FIELDS
            ),
            "family_provider_requirements": family_requirements,
            "promotion_status": "REQUESTED_PROVIDER_PACKET",
            "frequency_binding_status": "WAITING_FOR_REVIEWED_FREQUENCY_AND_EXACT_MATCH",
            "frequency_assigned": False,
            "release_values_present": False,
        })
    return {
        "schema": PROVIDER_PACKET_REQUEST_CATALOG_SCHEMA,
        "status": "PROVIDER_PACKET_REQUESTS_ONLY",
        "source_audit_schema": audit.get("schema"),
        "source_root": audit.get("source_root"),
        "source_audit_digest": hashlib.sha256(
            json.dumps(audit, ensure_ascii=False, sort_keys=True, separators=(",", ":")).encode("utf-8")
        ).hexdigest(),
        "request_count": len(requests),
        "candidate_family_request_counts": dict(sorted(family_counts.items())),
        "requests": requests,
        "provider_promotion": {
            "qualified_packet_count": 0,
            "frequency_assigned": False,
            "scenario_rows_are_not_release_packets": True,
            "next_input_schema": "prism.resolved_release.v1",
        },
        "blocking_reasons": [
            "scenario records are candidate hazards, not resolved release histories",
            "provider packet values and ownership remain missing until the external equipment model exports them",
        ],
    }


# ---------------------------------------------------------------------------
# Generic external CSV screening import
# ---------------------------------------------------------------------------

EXTERNAL_SCENARIO_EVIDENCE_SCHEMA = "prism.external_scenario_evidence.v1"

_EXTERNAL_CATEGORY_FAMILY_MAP: dict[str, tuple[str, ...]] = {
    "flange_connect": ("pipe_hole", "valve_hole"),
    "purge_vent": ("vent_hole", "vertical_handoff"),
    "valve_close": ("valve_hole",),
    "main_line_open": ("pipe_hole", "valve_hole"),
    "gas_phase_fill": ("vessel_hole", "pipe_hole"),
    "pressure_rise_monitor": ("vent_hole", "inventory_over_600s"),
    "relief_vent": ("vent_hole", "inventory_over_600s"),
    "supply_start": ("pipe_hole", "valve_hole"),
    "vaporizer_pass": ("vaporizer_hole",),
    "test_building_supply_gas": ("pipe_hole",),
}

_EXTERNAL_CATEGORY_SIGNAL_MAP: dict[str, tuple[str, ...]] = {
    "flange_connect": ("leak_test", "hydrogen_detection"),
    "purge_vent": ("vent_state", "hydrogen_detection", "pressure"),
    "valve_close": ("valve_state", "pressure", "mass_flow"),
    "main_line_open": ("valve_state", "pressure", "mass_flow"),
    "gas_phase_fill": ("valve_state", "pressure", "temperature"),
    "pressure_rise_monitor": ("pressure", "vacuum_level", "relief_state"),
    "relief_vent": ("pressure", "relief_state", "hydrogen_detection"),
    "supply_start": ("pressure", "mass_flow", "valve_state"),
    "vaporizer_pass": ("temperature", "mass_flow", "pressure"),
    "test_building_supply_gas": ("hydrogen_detection", "ventilation_state", "pressure"),
}

_EXTERNAL_FIELD_ALIASES: dict[str, tuple[str, ...]] = {
    "scenario": ("시나리오", "scenario"),
    "operating_mode": ("운전모드", "운전 모드", "operating_mode", "operation_mode", "mode"),
    "equipment": ("탱크", "tank", "공급모드", "supply_mode"),
    "step": ("절차번호", "step", "step_no"),
    "action": ("동작내용", "action", "operation"),
    "category": ("위험성분류", "risk_classification", "risk_class"),
    "leak_failure": ("누출_예상고장", "leak_failure", "release_failure"),
    "leak_risk": ("누출_위험", "leak_risk", "release_risk"),
    "leak_control": ("누출_조치방법", "leak_control", "release_control"),
    "leak_basis": ("누출_근거기준", "leak_basis", "release_basis"),
    "fire_risk": ("화재_위험", "fire_risk"),
    "explosion_risk": ("폭발_위험", "explosion_risk"),
}

def _external_value(row: Mapping[str, Any], aliases: tuple[str, ...]) -> str | None:
    for alias in aliases:
        value = row.get(alias)
        if value is not None and str(value).strip():
            return str(value).strip()
    return None


def _external_contains_any(text: str, terms: tuple[str, ...]) -> bool:
    lowered = text.casefold()
    return any(term.casefold() in lowered for term in terms)


def _external_scope_status(row_text: str) -> str:
    if _external_contains_any(row_text, ("밀폐", "반밀폐", "시험동", "건물", "indoor", "enclosed")):
        return "OUTSIDE_DECLARED_OUTDOOR_SCOPE"
    return "CANDIDATE_OUTDOOR_SCREENING"


def _external_signals(category: str | None, row_text: str) -> tuple[str, ...]:
    selected = list(_EXTERNAL_CATEGORY_SIGNAL_MAP.get(category or "", ()))
    text_rules = (
        ("pressure", ("압력", "pressure", "배압", "과압")),
        ("temperature", ("온도", "냉각", "temperature", "결빙")),
        ("mass_flow", ("유량", "flow", "공급")),
        ("valve_state", ("밸브", "XV", "valve", "개방", "폐쇄")),
        ("hydrogen_detection", ("가스검지", "검지기", "hydrogen detector", "누출감지")),
        ("ventilation_state", ("환기", "ventilation")),
        ("relief_state", ("PSV", "안전밸브", "relief", "벤트")),
        ("vacuum_level", ("진공", "vacuum")),
        ("leak_test", ("기밀", "누설시험", "leak test")),
    )
    for signal, terms in text_rules:
        if _external_contains_any(row_text, terms):
            selected.append(signal)
    return tuple(dict.fromkeys(selected))


def _external_record(
    row: Mapping[str, Any],
    *,
    source_path: Path,
    relative_path: str,
    source_sha256: str,
    row_number: int,
) -> dict[str, Any]:
    category = _external_value(row, _EXTERNAL_FIELD_ALIASES["category"])
    text = " ".join(str(value) for value in row.values() if value is not None)
    scenario = _external_value(row, _EXTERNAL_FIELD_ALIASES["scenario"])
    operating_mode_candidate, operating_mode_source = _operating_mode_candidate(
        _external_value(row, _EXTERNAL_FIELD_ALIASES.get("operating_mode", ())),
        scenario,
        relative_path,
    )
    return {
        "id": f"{relative_path}:{row_number}",
        "source_file": relative_path,
        "source_path": str(source_path),
        "source_sha256": source_sha256,
        "row_number": row_number,
        "scenario": scenario,
        "operating_mode_candidate": operating_mode_candidate,
        "operating_mode_source": operating_mode_source,
        "operating_mode_is_runtime_input": False,
        "equipment_tag": _external_value(row, _EXTERNAL_FIELD_ALIASES["equipment"]),
        "procedure_step": _external_value(row, _EXTERNAL_FIELD_ALIASES["step"]),
        "action": _external_value(row, _EXTERNAL_FIELD_ALIASES["action"]),
        "risk_classification": category,
        "candidate_accident_families": list(_EXTERNAL_CATEGORY_FAMILY_MAP.get(category or "", ())),
        "candidate_monitoring_signals": list(_external_signals(category, text)),
        "scope_status": _external_scope_status(text),
        "leak_failure": _external_value(row, _EXTERNAL_FIELD_ALIASES["leak_failure"]),
        "leak_risk": _external_value(row, _EXTERNAL_FIELD_ALIASES["leak_risk"]),
        "leak_control": _external_value(row, _EXTERNAL_FIELD_ALIASES["leak_control"]),
        "leak_basis": _external_value(row, _EXTERNAL_FIELD_ALIASES["leak_basis"]),
        "fire_risk": _external_value(row, _EXTERNAL_FIELD_ALIASES["fire_risk"]),
        "explosion_risk": _external_value(row, _EXTERNAL_FIELD_ALIASES["explosion_risk"]),
        "frequency_applied": False,
        "provider_packet_created": False,
        "facility_qra_claim": False,
    }


def import_scenario_evidence(root: str | Path) -> dict[str, Any]:
    """Import external procedure/failure CSVs as non-promoting evidence."""
    source_root = Path(root).resolve()
    if not source_root.is_dir():
        raise ValueError(f"scenario evidence root is not a directory: {source_root}")
    files = sorted(source_root.rglob("*.csv"))
    if not files:
        raise ValueError(f"no CSV scenario evidence found below {source_root}")
    records: list[dict[str, Any]] = []
    file_rows: list[dict[str, Any]] = []
    errors: list[str] = []
    for path in files:
        relative = path.relative_to(source_root).as_posix()
        digest = _sha256(path)
        row_count = 0
        try:
            with path.open("r", encoding="utf-8-sig", newline="") as handle:
                reader = csv.DictReader(handle)
                if reader.fieldnames is None:
                    raise ValueError("CSV header is missing")
                for row_number, row in enumerate(reader, start=2):
                    if not any(value is not None and str(value).strip() for value in row.values()):
                        continue
                    row_count += 1
                    records.append(_external_record(
                        row,
                        source_path=path,
                        relative_path=relative,
                        source_sha256=digest,
                        row_number=row_number,
                    ))
        except (OSError, UnicodeDecodeError, csv.Error, ValueError) as error:
            errors.append(f"{relative}: {error}")
        file_rows.append({
            "path": relative,
            "absolute_path": str(path),
            "sha256": digest,
            "row_count": row_count,
        })
    family_counts = Counter(
        family for record in records for family in record["candidate_accident_families"]
    )
    signal_counts = Counter(
        signal for record in records for signal in record["candidate_monitoring_signals"]
    )
    scope_counts = Counter(record["scope_status"] for record in records)
    operating_mode_counts = Counter(
        record["operating_mode_candidate"]
        for record in records
        if record.get("operating_mode_candidate") is not None
    )
    status = "SCREENING_EVIDENCE_IMPORTED" if records and not errors else (
        "PARTIAL_SCREENING_EVIDENCE" if records else "FAILED_SCREENING_EVIDENCE"
    )
    return {
        "schema_version": EXTERNAL_SCENARIO_EVIDENCE_SCHEMA,
        "status": status,
        "source_root": str(source_root),
        "source_file_count": len(files),
        "source_files": file_rows,
        "records": records,
        "summary": {
            "record_count": len(records),
            "family_candidate_counts": dict(sorted(family_counts.items())),
            "monitoring_signal_candidate_counts": dict(sorted(signal_counts.items())),
            "scope_status_counts": dict(sorted(scope_counts.items())),
            "operating_mode_candidate_counts": dict(sorted(operating_mode_counts.items())),
        },
        "errors": errors,
        "policy": {
            "facility_neutral": True,
            "frequency_applied": False,
            "sensor_frequency_learning_used": False,
            "provider_packet_created": False,
            "facility_qra_claim": False,
            "indoor_and_enclosed_rows_excluded_from_outdoor_scope": True,
            "candidate_family_map_is_not_a_frequency_map": True,
            "procedure_label_is_not_a_runtime_operating_mode": True,
        },
        "limitations": [
            "Procedure and failure text identifies candidate scenarios only; it does not provide an initiating frequency.",
            "Monitoring signal names are candidate tags and are not inverse leak-size or frequency estimators.",
            "Each candidate still requires an equipment-owned accident history, termination ledger and provider packet.",
            "Indoor, semi-enclosed and test-building rows are retained as evidence but marked outside the declared outdoor scope.",
        ],
    }
